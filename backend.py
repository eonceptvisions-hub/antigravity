import pandas as pd
from google import genai
from google.genai import types
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import time
import re
import json
import io
import base64
from pypdf import PdfReader
from docx import Document
from typing import Dict, Any, List
import os
import google.oauth2.credentials
from google_auth_oauthlib.flow import Flow
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError

# --- Constants & Config ---
LEAD_SCHEMA = ["id", "name", "company", "raw_source", "website", "mode", "contact_email", "draft_status", "send_status", "address", "profession", "phone", "draft_body"]

FUZZY_MAPPINGS = {
    "name": ["name", "person", "contact", "lead name", "owner", "founder"],
    "company": ["company", "firm", "business", "org", "organization", "biz"],
    "website": ["website", "url", "link", "site"],
    "contact_email": ["email", "e-mail", "mail", "contact_email"],
    "address": ["address", "location", "city", "area"],
    "profession": ["profession", "title", "role", "job"],
    "phone": ["phone", "mobile", "contact_number"]
}

def extract_text_from_file(file) -> str:
    """Extracts text from PDF, DOCX, or TXT files."""
    try:
        if file.name.lower().endswith('.pdf'):
            reader = PdfReader(file)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
            return text.strip()
        elif file.name.lower().endswith('.docx'):
            doc = Document(file)
            text = "\n".join([para.text for para in doc.paragraphs])
            return text.strip()
        elif file.name.lower().endswith('.txt'):
            return file.getvalue().decode("utf-8")
        else:
            raise ValueError(f"Unsupported format: {file.name}")
    except Exception as e:
        raise Exception(f"Extraction failed for {file.name}: {str(e)}")

def fuzzy_match_headers(csv_cols: List[str]) -> Dict[str, str]:
    """Maps CSV headers to our standardized SCHEMA using fuzzy keywords, avoiding duplicates."""
    mapping = {}
    used_targets = set()
    
    for col in csv_cols:
        clean_col = col.lower().strip()
        for target, keywords in FUZZY_MAPPINGS.items():
            if target in used_targets:
                continue
            if any(k == clean_col for k in keywords) or any(k in clean_col for k in keywords):
                mapping[col] = target
                used_targets.add(target)
                break
    return mapping

def process_upload(file, mode: str = "BUSINESS_SALES") -> pd.DataFrame:
    """Reads a CSV file, normalizes it to the Lead Schema, and drops invalid rows."""
    try:
        df = pd.read_csv(file)
    except Exception as e:
        raise ValueError(f"Error reading CSV file: {str(e)}")

    if df.empty:
        raise ValueError("The uploaded CSV file is empty.")

    # Standardize: Lowercase all existing headers first to catch simple case duplicates
    df.columns = [c.lower().strip() for c in df.columns]
    
    # Apply Fuzzy Mapping
    header_map = fuzzy_match_headers(df.columns)
    df = df.rename(columns=header_map)
    
    # Consolidate duplicate columns if any (e.g., if Name and Person both existed)
    df = df.loc[:, ~df.columns.duplicated()].copy()

    # Initialize Missing Schema Fields
    for col in LEAD_SCHEMA:
        if col not in df.columns:
            df[col] = ""
    
    # Fill defaults
    df["id"] = [str(uuid.uuid4()) for _ in range(len(df))]
    df["mode"] = mode
    df["draft_status"] = "PENDING"
    df["send_status"] = "QUEUED"
    df["raw_source"] = f"Upload: {file.name}"

    # Validation: Drop rows without Name AND Company
    df = df.dropna(subset=["name", "company"], how='all')
    # Filter out completely empty names
    df = df[df["name"].astype(str).str.strip() != ""]
    
    return df[LEAD_SCHEMA]

def predict_email_logic(website: str, mode: str) -> str:
    """Implements dual-mode email prediction strategy."""
    if not website or pd.isna(website):
        return ""
    
    domain = re.sub(r'https?://(www\.)?', '', website).split('/')[0]
    if not domain:
        return ""

    if mode == "JOB_HUNTING":
        # Strategy for hiring managers
        return f"careers@{domain}" # Basic fallback, LLM can refine
    else:
        # Strategy for business sales
        return f"info@{domain}"

def enrich_leads(df: pd.DataFrame) -> pd.DataFrame:
    """Predicts contact emails based on User Intent (mode)."""
    if "contact_email" not in df.columns:
        df["contact_email"] = ""

    mask = (df["contact_email"] == "") | (df["contact_email"].isna())
    df.loc[mask, "contact_email"] = df[mask].apply(
        lambda x: predict_email_logic(x["website"], x["mode"]), axis=1
    )
    return df

def generate_search_prompt(query: str) -> str:
    """Expert OSINT Research Protocol."""
    return f"""
    You are an expert Lead Generation Specialist and Open Source Intelligence (OSINT) analyst. Your goal is to compile a high-quality, verified database of leads related to: "{query}".

    ### MISSION OBJECTIVE
    Generate a clean, verified JSON dataset of 50-100 unique leads. Do not hallucinate data.

    ### RESEARCH PROTOCOL
    1. Identify the Entity/Decision Maker.
    2. Cross-reference presence in Gurgaon (if applicable) or target sector.
    3. Prioritize direct emails over generic.

    ### STRICT OUTPUT SCHEMA
    Return ONLY a valid JSON list. Fields: name, company, address, profession, website, contact_email, phone.
    """

def get_strategist_response(client, messages, global_context) -> str:
    """Chatbot logic with SaaS Context grounding."""
    system_instruction = f"""
    You are the 'Outreach Strategic Brain'. Your goal is to help refine outreach payloads.
    Current Global Context: {global_context}
    """
    try:
        response = client.models.generate_content(
            model='gemini-2.0-flash-exp',
            contents=messages,
            config=types.GenerateContentConfig(system_instruction=system_instruction)
        )
        return response.text
    except Exception as e:
        return f"Brain Error: {str(e)}"

def search_and_extract_leads(query: str, api_key: str, mode: str, custom_prompt: str = None) -> pd.DataFrame:
    """Lead Finder using Search Grounding and SaaS normalization."""
    client = genai.Client(api_key=api_key)
    prompt = custom_prompt if custom_prompt else generate_search_prompt(query)

    try:
        response = client.models.generate_content(
            model='gemini-2.0-flash-exp',
            contents=prompt,
            config=types.GenerateContentConfig(tools=[types.Tool(google_search=types.GoogleSearch())])
        )
        
        json_match = re.search(r'\[.*\]', response.text, re.DOTALL)
        if json_match:
            data = json.loads(json_match.group(0))
            df = pd.DataFrame(data)
            
            # Standardize: Lowercase headers
            df.columns = [c.lower().strip() for c in df.columns]

            # Normalize to Schema
            header_map = fuzzy_match_headers(df.columns)
            df = df.rename(columns=header_map)
            
            # Consolidate duplicates
            df = df.loc[:, ~df.columns.duplicated()].copy()
            
            for col in LEAD_SCHEMA:
                if col not in df.columns: df[col] = ""
            
            df["id"] = [str(uuid.uuid4()) for _ in range(len(df))]
            df["mode"] = mode
            df["draft_status"] = "PENDING"
            df["send_status"] = "QUEUED"
            df["raw_source"] = "Search Grounding"
            
            return df[LEAD_SCHEMA]
        raise ValueError("No JSON found")
    except Exception as e:
        raise Exception(f"Search failed: {str(e)}")

def generate_email_drafts(df: pd.DataFrame, api_key: str, sender_context: str) -> pd.DataFrame:
    """LLM Agent for personalized drafting using SaaS attributes."""
    client = genai.Client(api_key=api_key)

    for index, row in df.iterrows():
        if row.get("draft_status") == "GENERATED": continue

        mode_instruction = "Focus on hiring managers and human resources." if row["mode"] == "JOB_HUNTING" else "Focus on business owners and decision makers."
        
        prompt = f"""
        Role: Senior Sales Development Representative (SDR)
        Target Person: {row['name']}
        Target Company: {row['company']}
        Target Profession: {row.get('profession', 'Professional')}
        Mode/Mission: {row['mode']}
        
        Sender's Value Proposition & Context:
        {sender_context}
        
        TASK: Write a highly personalized, low-friction cold email. 
        
        STRICT GUIDELINES:
        1. Keep it under 100 words.
        2. No generic "I hope this email finds you well" or "My name is...".
        3. Hook them with a relevant observation related to their company/role.
        4. Focus on ONE specific problem we solve for them.
        5. End with a soft call-to-action (e.g., "Open to a brief chat next week?").
        6. Tone: Professional but conversational and direct.
        
        OUTPUT: Return ONLY the email body.
        """

        try:
            response = client.models.generate_content(
                model='gemini-2.0-flash-exp',
                contents=prompt,
                config=types.GenerateContentConfig(tools=[types.Tool(google_search=types.GoogleSearch())])
            )
            df.at[index, "draft_body"] = response.text
            df.at[index, "draft_status"] = "GENERATED"
            time.sleep(0.5)
        except:
            df.at[index, "draft_status"] = "FAILED"

    return df

def is_missing(value) -> bool:
    """Checks if a value is a placeholder or effectively empty."""
    if pd.isna(value): return True
    s = str(value).strip().lower()
    return s in ["", "0", "nan", "none", "n/a", "[not_found]"]

def fill_missing_info(row: Dict[str, Any], api_key: str) -> Dict[str, str]:
    """
    Uses Gemini OSINT protocol to find missing info for a single lead.
    Recognizes placeholders like '0' as missing data.
    """
    client = genai.Client(api_key=api_key)
    name = row.get("name", "Unknown")
    company = row.get("company", "Unknown")
    mode = row.get("mode", "BUSINESS_SALES")
    
    # Identify fields that are truly missing or contain placeholders
    target_fields = ["website", "contact_email", "phone", "address"]
    missing_fields = [f for f in target_fields if is_missing(row.get(f))]
    
    if not missing_fields:
        return {}

    # Mission-specific targeting logic
    target_description = ""
    if mode == "JOB_HUNTING":
        target_description = "Specifically look for recruiting emails, HR contacts, or their professional personal email address."
    else:
        target_description = "Specifically look for contact emails, business owner emails, or general 'connect' emails for the team."

    prompt = f"""
    You are an expert OSINT researcher. Your mission is to find missing business data for:
    ENTITY: {name}
    COMPANY: {company}
    
    SEARCH GOALS: Find values for {', '.join(missing_fields)}.
    Note: These fields are currently marked with '0' or placeholders in our vault, indicating they need verification.
    
    MISSION CONTEXT: {mode}
    STRATEGY: {target_description}
    
    STEPS:
    1. Use Google Search to find official corporate sites or high-authority directory listings (LinkedIn, Crunchbase, etc.).
    2. Extract and verify the specific {', '.join(missing_fields)}.
    3. If previous searches ([NOT_FOUND]) are mentioned, try a more specific search (e.g., search for the person's name directly with the company).
    
    Return ONLY a valid JSON object with these keys: {missing_fields}
    If you cannot find a specific field, return it as an empty string.
    """

    try:
        response = client.models.generate_content(
            model='gemini-2.0-flash-exp',
            contents=prompt,
            config=types.GenerateContentConfig(
                tools=[types.Tool(google_search=types.GoogleSearch())]
            )
        )
        
        # Debug print (safely)
        try:
            print(f"--- Raw Response for {name} ---\n{response.text}\n----------------")
        except:
            pass

        json_match = re.search(r'\{.*\}', response.text, re.DOTALL)
        if json_match:
            raw_json = json.loads(json_match.group(0))
            normalized_json = {k.lower(): v for k, v in raw_json.items()}
            return {field: normalized_json.get(field.lower(), "") for field in missing_fields}
    except Exception as e:
        try:
            print(f"OSINT Error for {name}: {e}")
        except:
            pass
    
    return {}

# --- Google OAuth2 & Gmail API ---

def get_google_auth_url(client_config: Dict[str, Any], redirect_uri: str) -> str:
    """Generates the Google OAuth2 authorization URL."""
    flow = Flow.from_client_config(
        client_config,
        scopes=['https://www.googleapis.com/auth/gmail.send'],
        redirect_uri=redirect_uri
    )
    auth_url, _ = flow.authorization_url(prompt='consent', access_type='offline')
    return auth_url

def get_google_credentials(client_config: Dict[str, Any], redirect_uri: str, code: str):
    """Exchanges the authorization code for credentials."""
    flow = Flow.from_client_config(
        client_config,
        scopes=['https://www.googleapis.com/auth/gmail.send'],
        redirect_uri=redirect_uri
    )
    flow.fetch_token(code=code)
    return flow.credentials

def send_email_via_gmail_api(credentials, recipient: str, subject: str, body: str):
    """Sends an email using the Gmail API."""
    try:
        service = build('gmail', 'v1', credentials=credentials)
        message = MIMEMultipart()
        message['To'] = recipient
        message['Subject'] = subject
        message.attach(MIMEText(body, 'plain'))
        
        raw_message = base64.urlsafe_b64encode(message.as_bytes()).decode()
        sent_message = service.users().messages().send(userId="me", body={'raw': raw_message}).execute()
        return sent_message
    except HttpError as error:
        raise Exception(f"Gmail API Error: {error}")
    except Exception as e:
        raise Exception(f"Dispatch Error: {str(e)}")

def send_batch_emails_api(df: pd.DataFrame, credentials) -> Dict[str, int]:
    """Dispatch logic via Gmail API with status updates."""
    results = {"sent": 0, "failed": 0}
    for index, row in df.iterrows():
        if row.get("send_status") == "SENT": continue
        if not row["contact_email"] or not row["draft_body"]:
            df.at[index, "send_status"] = "FAILED"
            results["failed"] += 1
            continue

        try:
            send_email_via_gmail_api(
                credentials, 
                row["contact_email"], 
                "Protocol: Outreach Connection", 
                row["draft_body"]
            )
            df.at[index, "send_status"] = "SENT"
            results["sent"] += 1
        except:
            df.at[index, "send_status"] = "FAILED"
            results["failed"] += 1
            
    return results

def send_batch_emails(df: pd.DataFrame, sender_email: str, app_password: str) -> Dict[str, int]:
    """Dispatch logic with status updates."""
    results = {"sent": 0, "failed": 0}
    try:
        server = smtplib.SMTP("smtp.gmail.com", 587)
        server.starttls()
        try:
            server.login(sender_email, app_password)
        except smtplib.SMTPAuthenticationError:
            raise Exception("Authentication Failed: Please ensure you are using a 16-character 'App Password' from Google, NOT your regular password. 2-Step Verification must be enabled.")

        for index, row in df.iterrows():
            if row.get("send_status") == "SENT": continue
            if not row["contact_email"] or not row["draft_body"]:
                df.at[index, "send_status"] = "FAILED"
                results["failed"] += 1
                continue

            try:
                msg = MIMEMultipart()
                msg['From'], msg['To'] = sender_email, row["contact_email"]
                msg['Subject'] = "Protocol: Outreach Connection"
                msg.attach(MIMEText(row["draft_body"], 'plain'))
                server.send_message(msg)
                df.at[index, "send_status"] = "SENT"
                results["sent"] += 1
            except:
                df.at[index, "send_status"] = "FAILED"
                results["failed"] += 1

        server.quit()
    except Exception as e:
        raise Exception(f"SMTP Error: {str(e)}")
    return results

def save_vault(df: pd.DataFrame, file_path: str = "vault.csv"):
    """Persists the vault DataFrame to a local CSV."""
    if df is not None:
        df.to_csv(file_path, index=False)

def load_vault(file_path: str = "vault.csv") -> pd.DataFrame:
    """Loads the vault from a local CSV if it exists."""
    import os
    if os.path.exists(file_path):
        try:
            return pd.read_csv(file_path)
        except:
            return None
    return None
